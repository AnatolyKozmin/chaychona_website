"""Разбор «Вкусной тетради» в Word: четыре вида вёрстки таблиц и фото из ячеек.

Документы собираем здесь же через python-docx: держать в репозитории 16-мегабайтный
файл шефа ради теста незачем, а вёрстку, из-за которой разбор и написан, проще
воспроизвести явно — видно, что именно проверяется.
"""
import io

import docx
import pytest
from docx.shared import Pt

from app.services.menu_import import RegistryParseError, parse_registry

# 1×1 PNG — картинке в тесте важно быть валидной, а не красивой.
PNG_1PX = bytes.fromhex(
    "89504e470d0a1a0a0000000d4948445200000001000000010802000000907753de"
    "0000000c49444154789c6338312d050003ec01c30a3511420000000049454e44ae426082"
)


def _doc_bytes(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _add_photo(cell, name: str) -> None:
    """Положить картинку в ячейку рядом с текстом — как в документе шефа."""
    cell.paragraphs[0].add_run().add_picture(io.BytesIO(PNG_1PX), width=Pt(20))


def test_plain_table_gives_dish_with_photo_and_texts():
    document = docx.Document()
    document.add_paragraph("Вкусная Тетрадь")
    document.add_paragraph("Восточная кухня")
    document.add_paragraph("Салаты")
    document.add_paragraph("Время приготовления 10-15 минут")
    table = document.add_table(rows=2, cols=3)
    table.rows[0].cells[0].text = "Блюдо"
    table.rows[0].cells[1].text = "Состав"
    table.rows[0].cells[2].text = "Красочное описание"
    body = table.rows[1].cells
    body[0].text = "Салат Айчичук"
    _add_photo(body[0], "айчичук")
    body[1].text = "Основные ингредиенты:\nТоматы, лук репчатый, базилик"
    body[2].text = "Яркий и свежий салат"

    registry = parse_registry(_doc_bytes(document), "тетрадь.docx")

    assert len(registry.rows) == 1
    row = registry.rows[0]
    assert row.name == "Салат Айчичук"
    assert row.branch == "Восточная кухня"
    assert row.category == "Салаты"
    # Подпись «Основные ингредиенты:» — оформление, в состав блюда её не тащим.
    assert row.ingredients == "Томаты, лук репчатый, базилик"
    assert row.description == "Яркий и свежий салат"
    assert registry.read_media(row.photo_dish) == PNG_1PX


def test_transposed_table_reads_names_above_and_composition_below():
    """Соусы свёрстаны поперёк: имена в первой строке, состав — во второй."""
    document = docx.Document()
    document.add_paragraph("Восточная кухня")
    document.add_paragraph("Соусы на выбор")
    table = document.add_table(rows=2, cols=2)
    for column, name in enumerate(("Сацебели", "Тайский")):
        table.rows[0].cells[column].text = name
        _add_photo(table.rows[0].cells[column], name)
    table.rows[1].cells[0].text = "томатная паста, чеснок, хмели-сунели"
    table.rows[1].cells[1].text = "соевый соус, свит-чили, майонез"

    rows = parse_registry(_doc_bytes(document), "тетрадь.docx").rows

    assert [row.name for row in rows] == ["Сацебели", "Тайский"]
    assert rows[0].ingredients == "томатная паста, чеснок, хмели-сунели"
    assert rows[1].ingredients == "соевый соус, свит-чили, майонез"
    # Нестандартная вёрстка — помечаем, чтобы залив можно было проверить глазами.
    assert all(row.note for row in rows)


def test_tile_table_splits_name_and_composition_inside_one_cell():
    """Суши и гарниры свёрстаны плиткой: в ячейке название, под ним состав."""
    document = docx.Document()
    document.add_paragraph("Восточная кухня")
    document.add_paragraph("Суши-Бар")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Калифорния\nКраб, огурец, авокадо, рис"
    table.rows[0].cells[1].text = "Филадельфия\nСемга, сливочный сыр, рис"

    rows = parse_registry(_doc_bytes(document), "тетрадь.docx").rows

    assert [row.name for row in rows] == ["Калифорния", "Филадельфия"]
    assert rows[0].ingredients == "Краб, огурец, авокадо, рис"
    assert rows[1].description is None


def test_names_only_table_still_creates_dishes_with_a_note():
    """Варенье в документе идёт списком имён — блюда нужны, но с пометкой."""
    document = docx.Document()
    document.add_paragraph("Мировая Кухня")
    document.add_paragraph("Варенье")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Варенье кизил"
    table.rows[0].cells[1].text = "Варенье абрикос"

    rows = parse_registry(_doc_bytes(document), "тетрадь.docx").rows

    assert [row.name for row in rows] == ["Варенье кизил", "Варенье абрикос"]
    assert all(row.ingredients is None for row in rows)
    assert all("нет ни состава" in (row.note or "") for row in rows)


def test_same_section_in_two_kitchens_stays_apart():
    """«Салаты» есть в обеих кухнях — это разные разделы, схлопывать нельзя."""
    document = docx.Document()
    document.add_paragraph("Восточная кухня")
    document.add_paragraph("Салаты")
    east = document.add_table(rows=1, cols=2)
    east.rows[0].cells[0].text = "Салат Ташкент"
    east.rows[0].cells[1].text = "Основные ингредиенты:\nтелятина, редька"
    document.add_paragraph("Мировая Кухня")
    document.add_paragraph("Салаты")
    world = document.add_table(rows=1, cols=2)
    world.rows[0].cells[0].text = "Салат Цезарь"
    world.rows[0].cells[1].text = "Основные ингредиенты:\nромано, курица, пармезан"

    rows = parse_registry(_doc_bytes(document), "тетрадь.docx").rows

    assert [(row.branch, row.category) for row in rows] == [
        ("Восточная кухня", "Салаты"),
        ("Мировая Кухня", "Салаты"),
    ]


def test_merged_columns_do_not_duplicate_composition():
    """«Мангал» свёрстан объединёнными ячейками — состав не должен задвоиться."""
    document = docx.Document()
    document.add_paragraph("Восточная кухня")
    document.add_paragraph("Мангал")
    table = document.add_table(rows=1, cols=3)
    cells = table.rows[0].cells
    cells[0].text = "Шашлык из семги"
    cells[1].text = "Основные ингредиенты:\nСемга, лук, лаваш"
    cells[2].text = "Основные ингредиенты:\nСемга, лук, лаваш"

    rows = parse_registry(_doc_bytes(document), "тетрадь.docx").rows

    assert len(rows) == 1
    assert rows[0].ingredients == "Семга, лук, лаваш"
    assert rows[0].description is None


def test_document_without_tables_is_rejected():
    document = docx.Document()
    document.add_paragraph("Вкусная Тетрадь")
    document.add_paragraph("Восточная кухня")

    with pytest.raises(RegistryParseError, match="ни одного блюда"):
        parse_registry(_doc_bytes(document), "тетрадь.docx")
