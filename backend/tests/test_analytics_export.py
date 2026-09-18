"""Выгрузка результатов тестов в Excel.

Первая часть — сборка книги без базы: что на каком листе и как посчитано.
Вторая — эндпоинт целиком: фильтры, номер попытки по всей истории и лист
«Не проходили» на живом Postgres (пропускается, если базы нет).
"""
import io
import uuid
from datetime import datetime

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.services.analytics_export import AnswerRow, AttemptRow, EmployeeRow, build_workbook


def _attempt(**overrides) -> AttemptRow:
    base = dict(
        attempt_id=1,
        user_id="u1",
        user_name="Анна Петрова",
        user_login="anna",
        restaurant="Жизнь Удалась",
        job_title="Официант",
        user_active=True,
        test_id=10,
        test_title="Стандарты сервиса",
        attempt_number=1,
        finished_at=datetime(2026, 9, 8, 21, 30),  # 00:30 9 сентября по Москве
        duration_seconds=125,
        total_questions=10,
        correct_answers=6,
    )
    base.update(overrides)
    return AttemptRow(**base)


def _sheet_rows(book, title: str) -> list[tuple]:
    return [row for row in book[title].iter_rows(min_row=2, values_only=True)]


def _open(content: bytes):
    return load_workbook(io.BytesIO(content))


def test_workbook_has_a_sheet_for_every_question_the_client_asked():
    book = _open(build_workbook([_attempt()], [], [], {"Жизнь Удалась": 1}, []))

    assert book.sheetnames == [
        "Попытки", "Сотрудники", "Рестораны", "Тесты", "Сложные вопросы", "Не проходили", "Параметры",
    ]


def test_attempt_row_shows_moscow_time_score_and_duration():
    rows = _sheet_rows(_open(build_workbook([_attempt()], [], [], {}, [])), "Попытки")

    finished, name, login, restaurant, job, status, test, number, correct, total, percent, duration = rows[0]
    assert finished == datetime(2026, 9, 9, 0, 30)
    assert (name, restaurant, status, number) == ("Анна Петрова", "Жизнь Удалась", "работает", 1)
    assert (correct, total, percent, duration) == (6, 10, 60.0, "2:05")


def test_employee_sheet_keeps_best_and_last_result_separately():
    """Пересдача: лучший и последний результат — разные числа, и нужны оба."""
    attempts = [
        _attempt(attempt_id=1, attempt_number=1, correct_answers=9, finished_at=datetime(2026, 9, 1, 10)),
        _attempt(attempt_id=2, attempt_number=2, correct_answers=5, finished_at=datetime(2026, 9, 2, 10)),
    ]

    row = _sheet_rows(_open(build_workbook(attempts, [], [], {}, [])), "Сотрудники")[0]

    assert row[5:8] == (2, 90.0, 50.0)


def test_restaurant_sheet_counts_who_never_took_a_test():
    attempts = [_attempt(user_id="u1"), _attempt(attempt_id=2, user_id="u2", correct_answers=8)]

    row = _sheet_rows(_open(build_workbook(attempts, [], [], {"Жизнь Удалась": 5}, [])), "Рестораны")[0]

    name, team, took, never, attempts_count, average = row
    assert (name, team, took, never, attempts_count) == ("Жизнь Удалась", 5, 2, 3, 2)
    assert average == 70.0


def test_hard_questions_are_sorted_by_error_rate():
    answers = [
        AnswerRow(1, "Стандарты сервиса", "Лёгкий", True),
        AnswerRow(1, "Стандарты сервиса", "Трудный", False),
        AnswerRow(2, "Стандарты сервиса", "Трудный", False),
        AnswerRow(2, "Стандарты сервиса", "Лёгкий", False),
    ]

    rows = _sheet_rows(_open(build_workbook([], answers, [], {}, [])), "Сложные вопросы")

    assert [(row[1], row[4]) for row in rows] == [("Трудный", 100.0), ("Лёгкий", 50.0)]


def test_people_without_attempts_are_listed():
    never = [EmployeeRow("u3", "Пётр Сидоров", "petr", "Жизнь Удалась", "Бармен")]

    rows = _sheet_rows(_open(build_workbook([], [], never, {}, [])), "Не проходили")

    assert rows == [("Пётр Сидоров", "petr", "Жизнь Удалась", "Бармен")]


# --- эндпоинт на живой базе


@pytest.fixture
def team(client: TestClient, auth_headers: dict) -> dict:
    marker = uuid.uuid4().hex[:8]
    restaurant = client.post(
        "/api/v1/users/catalog/restaurants", json={"name": f"__pytest_export_{marker}"}, headers=auth_headers
    ).json()
    job = client.post(
        "/api/v1/users/catalog/job-titles",
        json={"name": "Официант", "restaurant_id": restaurant["id"]},
        headers=auth_headers,
    ).json()
    test = client.post(
        "/api/v1/tests",
        json={
            "title": f"Выгрузка {marker}",
            "restaurant_id": restaurant["id"],
            "job_title_id": job["id"],
            "questions": [
                {
                    "text": "Сколько секунд на встречу гостя?",
                    "question_type": "single",
                    "options": [{"text": "30", "is_correct": True}, {"text": "300", "is_correct": False}],
                }
            ],
        },
        headers=auth_headers,
    ).json()

    def register(login: str) -> dict:
        resp = client.post(
            "/api/v1/auth/register",
            json={
                "first_name": "Тест",
                "last_name": login,
                "restaurant": restaurant["name"],
                "job_title": "Официант",
                "desired_login": login,
                "password": "пароль1",
            },
        )
        assert resp.status_code == 201, resp.text
        return {"Authorization": f"Bearer {resp.json()['access_token']}"}

    return {
        "restaurant": restaurant["name"],
        "test_id": test["id"],
        "active": register(f"active{marker}"),
        "idle_login": f"idle{marker}",
        "idle": register(f"idle{marker}"),
    }


def _take_and_submit(client: TestClient, headers: dict, test_id: int, correct: bool) -> None:
    taken = client.get(f"/api/v1/tests/{test_id}/take", headers=headers).json()
    question = taken["questions"][0]
    option = next(o for o in question["options"] if (o["text"] == "30") == correct)
    resp = client.post(
        f"/api/v1/tests/{test_id}/submit",
        json={"answers": [{"question_id": question["id"], "option_ids": [option["id"]]}]},
        headers=headers,
    )
    assert resp.status_code == 200, resp.text


def test_export_for_one_restaurant(client, auth_headers, team):
    _take_and_submit(client, team["active"], team["test_id"], correct=False)
    _take_and_submit(client, team["active"], team["test_id"], correct=True)

    resp = client.get(
        "/api/v1/tests/analytics/export", params={"restaurant": team["restaurant"]}, headers=auth_headers
    )

    assert resp.status_code == 200, resp.text
    assert "attachment" in resp.headers["content-disposition"]
    book = _open(resp.content)
    attempts = _sheet_rows(book, "Попытки")
    # Свежие сверху: вторая попытка первой строкой, номер считан по истории.
    assert [(row[7], row[10]) for row in attempts] == [(2, 100.0), (1, 0.0)]
    assert all(row[3] == team["restaurant"] for row in attempts)
    never = [row[1] for row in _sheet_rows(book, "Не проходили")]
    assert never == [team["idle_login"]]
    restaurant_row = _sheet_rows(book, "Рестораны")[0]
    assert restaurant_row[:4] == (team["restaurant"], 2, 1, 1)
    hard = _sheet_rows(book, "Сложные вопросы")[0]
    assert (hard[2], hard[3]) == (2, 1)


def test_export_date_filter_is_inclusive_moscow_days(client, auth_headers, team):
    _take_and_submit(client, team["active"], team["test_id"], correct=True)

    past = client.get(
        "/api/v1/tests/analytics/export",
        params={"restaurant": team["restaurant"], "date_from": "2020-01-01", "date_to": "2020-01-31"},
        headers=auth_headers,
    )
    wide = client.get(
        "/api/v1/tests/analytics/export",
        params={"restaurant": team["restaurant"], "date_from": "2020-01-01", "date_to": "2099-12-31"},
        headers=auth_headers,
    )

    assert _sheet_rows(_open(past.content), "Попытки") == []
    assert len(_sheet_rows(_open(wide.content), "Попытки")) == 1


def test_export_is_for_superadmin_only(client, team):
    resp = client.get("/api/v1/tests/analytics/export", headers=team["active"])

    assert resp.status_code == 403
