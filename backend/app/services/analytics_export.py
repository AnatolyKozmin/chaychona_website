"""Выгрузка результатов тестов в Excel.

Заказчику нужна картина по ресторану и команде, которую можно отфильтровать,
пересортировать и переслать управляющему, — то есть файл, а не страница.
Поэтому книга собирается из уже отфильтрованных строк, и каждый лист отвечает
на свой вопрос:

- «Попытки» — все прохождения построчно, с номером попытки сотрудника;
- «Сотрудники» — по человеку и тесту: сколько раз сдавал, лучший и последний результат;
- «Рестораны» — сколько людей в команде, сколько прошли хотя бы один тест;
- «Тесты» — средний результат по каждому тесту;
- «Сложные вопросы» — где ошибаются чаще всего;
- «Не проходили» — работающие сотрудники без единой попытки за период.

Сборка книги не ходит в базу: на вход — готовые строки, на выход — байты.
Так её можно проверить тестом без Postgres, а запросы остаются в эндпоинте.
"""
from __future__ import annotations

import io
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Время в базе — UTC без зоны. Рестораны живут по Москве, и «сдал в 23:40»
# не должно превращаться в «сдал в 20:40». Смещение фиксированное: переходов
# на летнее время в Москве нет с 2014 года, а базы часовых поясов в slim-образе
# может и не быть.
LOCAL_TZ = timezone(timedelta(hours=3), "МСК")

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="111827")
_GOOD_FILL = PatternFill("solid", fgColor="DCFCE7")
_MID_FILL = PatternFill("solid", fgColor="FEF9C3")
_LOW_FILL = PatternFill("solid", fgColor="FEE2E2")


@dataclass(frozen=True)
class AttemptRow:
    attempt_id: int
    user_id: str
    user_name: str
    user_login: str
    restaurant: str | None
    job_title: str | None
    user_active: bool
    test_id: int
    test_title: str
    attempt_number: int  # какая по счёту попытка этого сотрудника по этому тесту
    finished_at: datetime
    duration_seconds: int | None
    total_questions: int
    correct_answers: int

    @property
    def percent(self) -> float:
        return round(self.correct_answers / self.total_questions * 100, 1) if self.total_questions else 0.0


@dataclass(frozen=True)
class AnswerRow:
    attempt_id: int
    test_title: str
    question_text: str
    is_correct: bool


@dataclass(frozen=True)
class EmployeeRow:
    user_id: str
    user_name: str
    user_login: str
    restaurant: str | None
    job_title: str | None


def local_time(value: datetime) -> datetime:
    """UTC из базы → московское время без зоны (Excel зоны не понимает)."""
    return value.replace(tzinfo=timezone.utc).astimezone(LOCAL_TZ).replace(tzinfo=None)


def _duration(seconds: int | None) -> str:
    if seconds is None:
        return ""
    minutes, rest = divmod(int(seconds), 60)
    return f"{minutes}:{rest:02d}"


def _score_fill(percent: float) -> PatternFill:
    # Те же пороги, что у цвета результата на странице теста.
    if percent >= 80:
        return _GOOD_FILL
    if percent >= 60:
        return _MID_FILL
    return _LOW_FILL


def _write_sheet(ws, headers: list[str], rows: list[list], *, percent_columns: tuple[int, ...] = ()) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
        for column in percent_columns:
            cell = ws.cell(row=ws.max_row, column=column)
            if isinstance(cell.value, (int, float)):
                cell.fill = _score_fill(cell.value)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    for index, header in enumerate(headers, start=1):
        values = [header] + [row[index - 1] for row in rows[:500]]
        width = max(len(str(value)) if value is not None else 0 for value in values)
        ws.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 8), 60)


def build_workbook(
    attempts: list[AttemptRow],
    answers: list[AnswerRow],
    employees_without_attempts: list[EmployeeRow],
    team_sizes: dict[str, int],
    filters_description: list[str],
) -> bytes:
    """Собрать книгу. `team_sizes` — работающих сотрудников по ресторану."""
    wb = Workbook()

    # --- Попытки
    ws = wb.active
    ws.title = "Попытки"
    _write_sheet(
        ws,
        [
            "Дата (МСК)", "Сотрудник", "Логин", "Ресторан", "Должность", "Статус",
            "Тест", "Попытка №", "Верно", "Вопросов", "Результат, %", "Время, мин:сек",
        ],
        [
            [
                local_time(a.finished_at), a.user_name, a.user_login, a.restaurant or "", a.job_title or "",
                "работает" if a.user_active else "архив", a.test_title, a.attempt_number,
                a.correct_answers, a.total_questions, a.percent, _duration(a.duration_seconds),
            ]
            for a in sorted(attempts, key=lambda item: item.finished_at, reverse=True)
        ],
        percent_columns=(11,),
    )
    for cell in ws["A"][1:]:
        cell.number_format = "DD.MM.YYYY HH:MM"

    # --- Сотрудники: человек × тест
    by_user_test: dict[tuple[str, int], list[AttemptRow]] = defaultdict(list)
    for a in attempts:
        by_user_test[(a.user_id, a.test_id)].append(a)
    employee_rows = []
    for group in by_user_test.values():
        group.sort(key=lambda item: item.finished_at)
        first, last = group[0], group[-1]
        employee_rows.append(
            [
                first.user_name, first.restaurant or "", first.job_title or "",
                "работает" if first.user_active else "архив", first.test_title, len(group),
                max(item.percent for item in group), last.percent, local_time(last.finished_at),
            ]
        )
    employee_rows.sort(key=lambda row: (row[1], row[0], row[4]))
    ws = wb.create_sheet("Сотрудники")
    _write_sheet(
        ws,
        ["Сотрудник", "Ресторан", "Должность", "Статус", "Тест", "Попыток", "Лучший, %", "Последний, %", "Последняя попытка (МСК)"],
        employee_rows,
        percent_columns=(7, 8),
    )
    for cell in ws["I"][1:]:
        cell.number_format = "DD.MM.YYYY HH:MM"

    # --- Рестораны
    per_restaurant: dict[str, list[AttemptRow]] = defaultdict(list)
    for a in attempts:
        per_restaurant[a.restaurant or "без ресторана"].append(a)
    restaurant_rows = []
    for name in sorted(set(per_restaurant) | set(team_sizes)):
        group = per_restaurant.get(name, [])
        active_people = {a.user_id for a in group if a.user_active}
        team = team_sizes.get(name, 0)
        # Средний — по последней попытке человека в каждом тесте: пересдачи
        # не должны тянуть среднее ни вверх, ни вниз.
        latest: dict[tuple[str, int], AttemptRow] = {}
        for a in group:
            key = (a.user_id, a.test_id)
            if key not in latest or a.finished_at > latest[key].finished_at:
                latest[key] = a
        average = round(sum(a.percent for a in latest.values()) / len(latest), 1) if latest else ""
        restaurant_rows.append(
            [name, team, len(active_people), max(team - len(active_people), 0), len(group), average]
        )
    ws = wb.create_sheet("Рестораны")
    _write_sheet(
        ws,
        ["Ресторан", "Работает сотрудников", "Проходили тесты", "Не проходили", "Попыток", "Средний результат, %"],
        restaurant_rows,
        percent_columns=(6,),
    )

    # --- Тесты
    per_test: dict[str, list[AttemptRow]] = defaultdict(list)
    for a in attempts:
        per_test[a.test_title].append(a)
    test_rows = []
    for title, group in sorted(per_test.items()):
        best: dict[str, float] = {}
        for a in group:
            best[a.user_id] = max(best.get(a.user_id, 0.0), a.percent)
        test_rows.append(
            [
                title, len(group), len(best),
                round(sum(a.percent for a in group) / len(group), 1),
                round(sum(best.values()) / len(best), 1),
            ]
        )
    ws = wb.create_sheet("Тесты")
    _write_sheet(
        ws,
        ["Тест", "Попыток", "Сотрудников", "Средний результат, %", "Средний лучший результат, %"],
        test_rows,
        percent_columns=(4, 5),
    )

    # --- Сложные вопросы
    per_question: dict[tuple[str, str], list[bool]] = defaultdict(list)
    for answer in answers:
        per_question[(answer.test_title, answer.question_text)].append(answer.is_correct)
    question_rows = []
    for (title, question), results in per_question.items():
        wrong = sum(1 for ok in results if not ok)
        question_rows.append([title, question, len(results), wrong, round(wrong / len(results) * 100, 1)])
    question_rows.sort(key=lambda row: (-row[4], -row[3]))
    ws = wb.create_sheet("Сложные вопросы")
    _write_sheet(ws, ["Тест", "Вопрос", "Ответов", "Ошибок", "Ошибок, %"], question_rows)
    ws.column_dimensions["B"].width = 80
    for cell in ws["B"][1:]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    # --- Не проходили
    ws = wb.create_sheet("Не проходили")
    _write_sheet(
        ws,
        ["Сотрудник", "Логин", "Ресторан", "Должность"],
        [
            [e.user_name, e.user_login, e.restaurant or "", e.job_title or ""]
            for e in sorted(employees_without_attempts, key=lambda item: (item.restaurant or "", item.user_name))
        ],
    )

    # --- Что выгружено: без этого через неделю не понять, какой срез в файле
    ws = wb.create_sheet("Параметры")
    ws.append(["Выгружено (МСК)", local_time(datetime.utcnow()).strftime("%d.%m.%Y %H:%M")])
    for line in filters_description:
        label, _, value = line.partition(": ")
        ws.append([label, value])
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 60

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
